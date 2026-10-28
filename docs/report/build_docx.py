"""Render docs/report/sonuc_raporu.md into the official TÜBİTAK 2209-A final-report template (.docx).

    uv run --with python-docx python docs/report/build_docx.py                 # repository copy (no receipts)
    uv run --with python-docx python docs/report/build_docx.py --submission    # docs/report/teslim/ copy + Ek-2

The template (docs/report/2209a_sonuc_raporu_formati.docx, published by TÜBİTAK) keeps its title box, the
general-information table and the signature table; the "Sonuç Raporu Formatı" heading list is replaced by the
report body. Appendices follow the signature block on new pages: Ek-1 (the forward decision report produced by
the system) and Ek-2 (expense receipts). The receipts contain personal data, so they are read from the local,
git-ignored folder docs/report/teslim/makbuz/ and only the --submission build embeds them.

Supported Markdown: the header key/value table, # / ## / ### headings, paragraphs, "* " bullets, pipe tables,
**bold** and images ``![caption](path relative to docs/report)``.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.image.image import Image as DocxImage
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "2209a_sonuc_raporu_formati.docx"
SOURCE = HERE / "sonuc_raporu.md"
FORWARD_APPENDIX = HERE / "ek_ileri_kosu_2026-09.md"
RECEIPTS = HERE / "teslim" / "makbuz"
REPO_URL = "https://github.com/vedattatli/bistmas-2209a"
FONT = "Arial"
USABLE_CM = 16.0
URL = re.compile(r"(https?://[^\s)]*[^\s).,;:])")
FIGURE_WIDTHS = {"sekil_mimari.png": 12.5}  # cm; the tall architecture diagram fits beside the agent table

# expense receipts for Ek-2: (file stem in RECEIPTS, label)
RECEIPT_FILES = [
    ("makbuz_2026-08-04", "04.08.2026 — 9.999,88 TL, sipariş no GPA.3368-3344-9472-70602"),
    ("makbuz_2026-09-04", "04.09.2026 — 9.999,88 TL, sipariş no GPA.3368-3344-9472-70602..0 (aylık yenileme)"),
    ("makbuz_2026-10-18", "18.10.2026 — 6.999,99 TL, sipariş no GPA.3361-7364-6447-12686"),
]

# paragraph layouts: (alignment, space before pt, space after pt, line spacing, left indent cm, first-line cm)
LAYOUT = {
    "body": (WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6, 1.15, 0, 0),
    "bullet": (WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 3, 1.15, 0.6, -0.4),
    "h1": (WD_ALIGN_PARAGRAPH.LEFT, 0, 10, 1.0, 0, 0),
    "h2": (WD_ALIGN_PARAGRAPH.LEFT, 14, 6, 1.0, 0, 0),
    "h3": (WD_ALIGN_PARAGRAPH.LEFT, 10, 4, 1.0, 0, 0),
    "figure": (WD_ALIGN_PARAGRAPH.CENTER, 6, 2, 1.0, 0, 0),
    "caption": (WD_ALIGN_PARAGRAPH.CENTER, 0, 10, 1.0, 0, 0),
    "small": (WD_ALIGN_PARAGRAPH.LEFT, 0, 2, 1.0, 0, 0),
    "ref": (WD_ALIGN_PARAGRAPH.LEFT, 0, 5, 1.1, 1.0, -1.0),  # reference list: hanging indent
}


def parse(md: str) -> tuple[dict[str, str], str, list[str]]:
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    title = re.search(r"\*\*PROJE BAŞLIĞI:\*\* \*\*(.+?)\*\*", md).group(1)
    info = dict(re.findall(r"^\| ([^|]+?) \| ([^|]+?) \|$", md.split("## 1.")[0], flags=re.M))
    body = md[md.index("## 1.") :].split("\n---\n")[0]
    return info, title, body.splitlines()


def add_runs(par: Paragraph, text: str, size: float = 11, bold: bool = False, italic: bool = False) -> None:
    """Add text with **bold** and *italic* spans; web addresses become clickable links."""
    pieces = []
    for k, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        for i, sub in enumerate(re.split(r"(?<![\w*])\*([^*\s][^*]*?)\*(?![\w*])", chunk)):
            for m, piece in enumerate(URL.split(sub.replace("`", ""))):
                pieces.append((piece, k % 2 == 1, i % 2 == 1, m % 2 == 1))
    for piece, b, it, is_link in pieces:
        if not piece:
            continue
        run = par.add_run(piece)
        run.bold = bold or b
        run.italic = italic or it
        run.font.name = FONT
        run.font.size = Pt(size)
        if is_link:
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), par.part.relate_to(piece, RT.HYPERLINK, is_external=True))
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x9A)
            run.font.underline = True
            link.append(run._r)
            par._p.append(link)


def style_par(par: Paragraph, kind: str, extra_before: float = 0) -> None:
    align, before, after, line, left, first = LAYOUT[kind]
    pf = par.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before + extra_before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.left_indent = Cm(left)
    pf.first_line_indent = Cm(first)
    if kind in ("h1", "h2", "h3", "figure"):
        pf.keep_with_next = True


def _child(el, tag: str, **attrs: str):
    child = el.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        el.append(child)
    for k, v in attrs.items():
        child.set(qn(k), v)
    return child


def column_widths(rows: list[list[str]]) -> list[float]:
    """Share the usable width by a mix of the longest and the mean cell length of each column.

    Every column first gets room for its longest word (so headers such as "Sharpe" never break inside a word);
    the columns that are wider than that give back width in proportion to their slack.
    """
    weights, floors = [], []
    for j in range(len(rows[0])):
        lens = [len(r[j]) for r in rows]
        weights.append(max(5.0, 0.6 * min(max(lens), 60) + 0.4 * sum(lens) / len(lens)))
        longest_word = max(len(w) for r in rows for w in re.split(r"[\s/]+", r[j].replace("**", "")))
        floors.append(max(1.3, 0.2 * longest_word + 0.45))  # ~0.2 cm per 9-pt bold character plus cell padding
    widths = [max(f, USABLE_CM * w / sum(weights)) for w, f in zip(weights, floors, strict=True)]
    excess = sum(widths) - USABLE_CM
    slack = [w - f for w, f in zip(widths, floors, strict=True)]
    if excess > 0 and sum(slack) > 0:
        widths = [w - excess * s / sum(slack) for w, s in zip(widths, slack, strict=True)]
    return [w * USABLE_CM / sum(widths) for w in widths]


def format_table(table, widths_cm: list[float]) -> None:
    """Thin grey grid, shaded header row repeated on page breaks, fixed column widths, rows kept whole."""
    tbl_pr = table._tbl.tblPr
    borders = _child(tbl_pr, "w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _child(borders, f"w:{side}", **{"w:val": "single", "w:sz": "4", "w:space": "0", "w:color": "8C8C8C"})
    _child(tbl_pr, "w:tblLayout", **{"w:type": "fixed"})
    table.autofit = False
    grid = table._tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), widths_cm, strict=False):
        col.set(qn("w:w"), str(int(w * 567)))
    for i, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        _child(tr_pr, "w:cantSplit", **{"w:val": "true"})
        if i == 0:
            _child(tr_pr, "w:tblHeader", **{"w:val": "true"})
        for cell, w in zip(row.cells, widths_cm, strict=False):
            cell.width = Cm(w)
            if i == 0:
                _child(cell._tc.get_or_add_tcPr(), "w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": "E6ECF5"})


class Writer:
    """Inserts paragraphs, tables and pictures one after another behind a cursor element of the document body."""

    def __init__(self, doc, cursor) -> None:
        self.doc = doc
        self.cursor = cursor
        self.after_table = 0.0

    def _new(self) -> Paragraph:
        p = OxmlElement("w:p")
        self.cursor.addnext(p)
        self.cursor = p
        return Paragraph(p, self.doc._body)

    def para(self, kind: str) -> Paragraph:
        par = self._new()
        style_par(par, kind, extra_before=self.after_table)
        self.after_table = 0.0
        return par

    def text(self, kind: str, text: str, size: float = 11, bold: bool = False, italic: bool = False) -> None:
        par = self.para(kind)
        if kind in ("body", "bullet") and re.search(r"https?://|[^\s]{40,}", text):
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT  # long links/hashes would stretch justified lines
        add_runs(par, text, size=size, bold=bold, italic=italic)

    def page_break(self) -> None:
        self._new().add_run().add_break(WD_BREAK.PAGE)

    def table(self, rows: list[list[str]], size: float = 9) -> None:
        rows = [r for r in rows if not re.fullmatch(r"[-| :]+", "|".join(r))]
        ncol = max(len(r) for r in rows)
        rows = [r + [""] * (ncol - len(r)) for r in rows]
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        for i, r in enumerate(rows):
            for j, val in enumerate(r):
                cp = t.cell(i, j).paragraphs[0]
                cp.paragraph_format.space_before = Pt(1.5)
                cp.paragraph_format.space_after = Pt(1.5)
                add_runs(cp, val, size=size, bold=i == 0)
        format_table(t, column_widths(rows))
        self.cursor.addnext(t._tbl)
        self.cursor = t._tbl
        self.after_table = 8.0

    def picture(self, path: Path, caption: str | None = None, width_cm: float = USABLE_CM, max_height_cm: float = 21.0):
        img = DocxImage.from_file(str(path))
        width_cm = min(width_cm, max_height_cm * img.px_width / img.px_height)
        self.para("figure").add_run().add_picture(str(path), width=Cm(width_cm))
        if caption:
            self.text("caption", caption, size=9, italic=True)

    def markdown(self, lines: list[str], shift: int = 0) -> None:
        """Render Markdown lines; ``shift`` demotes headings (used for appendices)."""
        buf: list[str] = []
        table: list[list[str]] = []

        in_refs = False

        def flush_par() -> None:
            if buf:
                text = " ".join(s.strip() for s in buf)
                self.text("bullet" if text.startswith("• ") else "ref" if in_refs else "body", text)
                buf.clear()

        for line in lines:
            if line.startswith("|"):
                flush_par()
                table.append([c.strip() for c in line.strip().strip("|").split("|")])
                continue
            if table:
                self.table(table)
                table = []
            if m := re.match(r"^(#{1,3}) (.+)", line):
                flush_par()
                level = min(len(m.group(1)) + shift, 3)
                in_refs = "Kaynakça" in m.group(2)
                self.text("h3" if level == 3 else "h2", m.group(2), size=11 if level == 3 else 12, bold=True)
            elif m := re.match(r"!\[(.*)\]\((.+)\)", line.strip()):
                flush_par()
                path = (HERE / m.group(2)).resolve()
                self.picture(path, m.group(1), width_cm=FIGURE_WIDTHS.get(path.name, USABLE_CM))
            elif line.startswith(("* ", "- ")):
                flush_par()
                buf.append("• " + line[2:])
            elif line.startswith("  ") and buf:
                buf.append(line)
            elif not line.strip():
                flush_par()
            else:
                buf.append(line)
        flush_par()
        if table:
            self.table(table)


def add_page_numbers(doc) -> None:
    """Centred "Sayfa n / N" footer."""
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        par = footer.paragraphs[0]
        for r in list(par.runs):
            r._r.getparent().remove(r._r)
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(par, "Sayfa ", size=9)
        for k, instr in enumerate(("PAGE", "NUMPAGES")):
            # complex field (begin / instruction / separate / cached value / end) so the runs keep their font
            for part in ("begin", "instr", "separate", "value", "end"):
                run = par.add_run("1" if part == "value" else None)
                run.font.name = FONT
                run.font.size = Pt(9)
                if part == "instr":
                    el = OxmlElement("w:instrText")
                    el.set(qn("xml:space"), "preserve")
                    el.text = f" {instr} "
                    run._r.append(el)
                elif part != "value":
                    el = OxmlElement("w:fldChar")
                    el.set(qn("w:fldCharType"), part)
                    run._r.append(el)
            if k == 0:
                add_runs(par, " / ", size=9)


def tr_upper(text: str) -> str:
    return text.replace("i", "İ").replace("ı", "I").upper()


def fill_title_box(doc, title: str, executor: str, advisor: str) -> None:
    executor, advisor = tr_upper(executor), tr_upper(advisor)  # the title page carries the names in capitals
    texts = list(doc.element.body.iter(qn("w:t")))
    for prev, t in zip([None, *texts], texts, strict=False):
        if t.text and t.text.strip() == "ADI:" and prev is not None and (prev.text or "").endswith("YÜRÜTÜCÜSÜNÜN"):
            t.text = f" ADI: {executor}"  # the template splits this label over two runs
    for t in doc.element.body.iter(qn("w:t")):
        if t.text and "(koyu harflerle yazılacak)" in t.text:
            t.text = t.text.replace("(koyu harflerle yazılacak)", title)
            run = t.getparent()
            rpr = run.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                run.insert(0, rpr)
            if rpr.find(qn("w:b")) is None:
                rpr.append(OxmlElement("w:b"))  # the template asks for the title in bold
        elif t.text and t.text.strip() == "PROJE YÜRÜTÜCÜSÜNÜN ADI:":
            t.text = f"PROJE YÜRÜTÜCÜSÜNÜN ADI: {executor}"
        elif t.text and t.text.strip() == "DANIŞMANININ ADI:":
            t.text = f" DANIŞMANININ ADI: {advisor}"
    # the frame has a fixed height and the long title wraps over several lines, which pushes the advisor line out
    # of the frame: drop half of the empty lines between "SONUÇ RAPORU" and the title
    box = next(doc.element.body.iter(qn("w:txbxContent")))
    paras = list(box.iter(qn("w:p")))
    label = ["".join(t.text or "" for t in p.iter(qn("w:t"))).strip() for p in paras]
    top = label.index("SONUÇ RAPORU")
    head = next(i for i, s in enumerate(label) if s.startswith("PROJE BAŞLIĞI"))
    gap = paras[top + 1 : head]
    for p in gap[: len(gap) // 2]:
        p.getparent().remove(p)


def build(out: Path, submission: bool) -> None:
    info, title, lines = parse(SOURCE.read_text(encoding="utf-8"))
    doc = docx.Document(str(TEMPLATE))
    executor = info.get("Proje yürütücüsünün adı", "").split(" (")[0]
    advisor = info.get("Danışmanın adı", "")
    fill_title_box(doc, title, executor, advisor)

    # a stray "ÜNÜ" run sits under the title box in the published template; drop it
    for t in list(doc.element.body.iter(qn("w:t"))):
        if t.text == "ÜNÜ":
            run = t.getparent()
            run.getparent().remove(run)

    # general-information table, then the remaining project details below it
    keys = {
        "PROJENİN KONUSU": "Projenin konusu",
        "PROJE YÜRÜTÜCÜSÜNÜN ADI": "Proje yürütücüsünün adı",
        "DANIŞMANIN ADI": "Danışmanın adı",
        "PROJE  BAŞLANGIÇ VE BİTİŞ TARİHLERİ": "Proje başlangıç ve bitiş tarihleri",
    }
    general = doc.tables[0]
    for row in general.rows:
        label = row.cells[0].text.strip()
        if label in keys:
            add_runs(row.cells[1].paragraphs[0], info.get(keys[label], ""), size=10)
    details = Writer(doc, general._tbl)
    details.after_table = 6.0
    details.text("small", f"**Başvuru numarası:** {info.get('Başvuru numarası', '')}", size=10)
    details.text("small", f"**Proje ekibi:** {info.get('Proje ekibi', '')}", size=10)
    details.text("small", "**Destek programı ve tutarı:** TÜBİTAK 2209-A, 9.000 TL", size=10)
    details.text("small", f"**Açık kaynak kod deposu:** {REPO_URL}", size=10)

    # signature table: names under the labels (signed by hand or e-signature), date line
    report_date = info.get("Proje başlangıç ve bitiş tarihleri", "").split("–")[-1].strip()
    sign = doc.tables[-1]
    for cell, name in zip(sign.rows[1].cells, (executor, advisor), strict=False):
        add_runs(cell.paragraphs[0], name, size=10)
    date_par = next(p for p in doc.paragraphs if p.text.strip().startswith("Tarih"))
    add_runs(date_par, f" {report_date}", size=11)
    # keep the signature table, the blank lines below it and the date line on one page
    for row in sign.rows:
        for cell in row.cells:
            for cp in cell.paragraphs:
                cp.paragraph_format.keep_with_next = True
    el = sign._tbl.getnext()
    while el is not None and el is not date_par._p:
        if el.tag == qn("w:p"):
            Paragraph(el, doc._body).paragraph_format.keep_with_next = True
        el = el.getnext()

    # replace the "Sonuç Raporu Formatı" heading list with the body
    paras = doc.paragraphs
    start = next(i for i, p in enumerate(paras) if p.text.strip().startswith("Sonuç Raporu Formatı"))
    end = next(i for i, p in enumerate(paras) if p.text.strip().startswith("Proje ile ilgili harcama"))
    anchor = paras[start]._p
    for p in paras[start + 1 : end + 1]:
        p._p.getparent().remove(p._p)
    Writer(doc, anchor).markdown(lines)
    anchor.getparent().remove(anchor)

    # appendices after the signature block, each on a new page
    ek = Writer(doc, date_par._p)
    ek.page_break()
    ek.text("h1", "EKLER", size=13, bold=True)
    ek.text("bullet", "• Ek-1. Eylül 2026 dönemi ileriye dönük karar raporu (sistem çıktısı)")
    ek.text("bullet", "• Ek-2. Harcama belgeleri (Google Play sipariş makbuzları)")
    ek.text("h2", "Ek-1. Eylül 2026 dönemi ileriye dönük karar raporu", size=12, bold=True)
    ek.text(
        "body",
        "Aşağıda, sistemin 30.09.2026 karar tarihi için ürettiği aylık karar raporu özgün hâliyle verilmiştir "
        "(yazılımın dili İngilizce olduğundan rapor İngilizcedir). Karar, ilgili ay kapandıktan sonra üretildiği için "
        '"geç" olarak etiketlenmiş ve ileriye dönük değerlendirmede ayrı tutulmuştur (bkz. 3.6).',
    )
    ek.markdown(FORWARD_APPENDIX.read_text(encoding="utf-8").splitlines(), shift=1)
    ek.page_break()
    ek.text("h2", "Ek-2. Harcama belgeleri", size=12, bold=True)
    if submission:
        ek.text(
            "body",
            "Proje desteği, proje boyunca yazılım geliştirme ve raporlamada kullanılan Anthropic Claude Max yapay zekâ "
            "hizmeti aboneliğine harcanmıştır (bkz. 6. bölüm). Abonelik Google Play üzerinden alınmış olup aşağıda "
            "yürütücü adına düzenlenmiş üç sipariş makbuzu yer almaktadır. Makbuzların ikinci sayfaları yalnızca "
            "satıcının genel iade koşullarını içerdiğinden eklenmemiştir.",
        )
        for k, (stem, label) in enumerate(RECEIPT_FILES, start=1):
            if k > 1:
                ek.page_break()
            ek.text("h3", f"Ek-2.{k}. {label}", size=10.5, bold=True)
            ek.picture(RECEIPTS / f"{stem}-1k.png", max_height_cm=18.5)
    else:
        ek.text(
            "body",
            "Harcama belgeleri kişisel bilgi içerdiğinden depodaki bu kopyada yer almamaktadır; e-BİDEB sistemine "
            "yüklenen nüshada bulunmaktadır. Harcama kalemleri 6. bölümde listelenmiştir.",
        )

    add_page_numbers(doc)
    doc.save(str(out))
    print("written", out)


if __name__ == "__main__":
    submission = "--submission" in sys.argv
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    default = HERE / "teslim" / "sonuc_raporu_teslim.docx" if submission else HERE / "sonuc_raporu.docx"
    out = Path(args[0]) if args else default
    out.parent.mkdir(parents=True, exist_ok=True)
    build(out, submission)
